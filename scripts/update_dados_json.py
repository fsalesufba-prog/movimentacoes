import json
import os
import argparse
from collections import OrderedDict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from zoneinfo import ZoneInfo

from openpyxl import load_workbook


SOURCE_DIR = r"C:\Users\MiltonOrricoVLTLote2\VLT Lote 2\VLT - LOTE 2 - 03 - Custos\06- Equipamentos"
SOURCE_FILE_HINT = "Moimentação Material-08-2025.xlsx"
SOURCE_SHEET = "Dados"
OUTPUT_FILE = os.path.join(os.path.dirname(os.path.dirname(__file__)), "dados.json")
TZ = ZoneInfo("America/Sao_Paulo")

HEADER_MAP = {
    "PATRIMÔNIO": "patrimonio",
    "EQUIPAMENTO": "equipamento",
    "PLACA": "placa_locador",
    "DATA": "data",
    "MATERIAL": "material",
    "ORIGEM": "origem",
    "DESTINO": "destino",
    "Nº VIAGENS": "num_viagens",
    "VOLUME": "volume_m3",
    "MODALIDADE": "modalidade",
    "VALOR": "valor",
}

OUTPUT_FIELD_ORDER = [
    "patrimonio",
    "equipamento",
    "placa_locador",
    "data",
    "material",
    "origem",
    "destino",
    "modalidade",
    "num_viagens",
    "volume_m3",
    "valor",
]


def find_source_file():
    for name in os.listdir(SOURCE_DIR):
        if name == SOURCE_FILE_HINT:
            return os.path.join(SOURCE_DIR, name)
    for name in os.listdir(SOURCE_DIR):
        if name.lower().endswith(".xlsx") and "08-2025" in name:
            return os.path.join(SOURCE_DIR, name)
    raise FileNotFoundError(f"Arquivo não encontrado em {SOURCE_DIR}")


def normalize_text(value):
    if value is None:
        return ""
    return str(value).strip()


def normalize_number(value):
    if value in (None, ""):
        return 0
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else float(value)
    text = str(value).strip().replace(".", "").replace(",", ".")
    try:
        number = Decimal(text)
    except InvalidOperation:
        return 0
    return int(number) if number == number.to_integral() else float(number)


def normalize_date(value):
    if value in (None, ""):
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    if " " in text:
        text = text.split(" ", 1)[0]
    return text


def read_rows(path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook[SOURCE_SHEET]
    rows = worksheet.iter_rows(values_only=True)

    raw_headers = next(rows)
    source_headers = [normalize_text(h).upper() for h in raw_headers]

    expected_headers = set(HEADER_MAP.keys())
    if set(source_headers) != expected_headers:
        raise ValueError(
            f"Cabeçalhos inesperados na aba {SOURCE_SHEET}: {source_headers}"
        )

    data = []
    for row in rows:
      if not row or all(cell in (None, "") for cell in row):
          continue

      source_row = dict(zip(source_headers, row))
      record = OrderedDict()
      for target_key in OUTPUT_FIELD_ORDER:
          source_key = next(key for key, value in HEADER_MAP.items() if value == target_key)
          raw_value = source_row.get(source_key)
          if target_key == "data":
              record[target_key] = normalize_date(raw_value)
          elif target_key in {"num_viagens", "volume_m3", "valor"}:
              record[target_key] = normalize_number(raw_value)
          else:
              record[target_key] = normalize_text(raw_value)
      data.append(record)

    return data


def read_rows_by_excel_range(path, start_row, end_row):
    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook[SOURCE_SHEET]

    source_headers = [
        normalize_text(worksheet.cell(1, col).value).upper()
        for col in range(1, worksheet.max_column + 1)
    ]
    expected_headers = set(HEADER_MAP.keys())
    if set(source_headers) != expected_headers:
        raise ValueError(
            f"Cabeçalhos inesperados na aba {SOURCE_SHEET}: {source_headers}"
        )

    rows = []
    for excel_row in range(start_row, end_row + 1):
        values = [worksheet.cell(excel_row, col).value for col in range(1, worksheet.max_column + 1)]
        if all(cell in (None, "") for cell in values):
            continue

        source_row = dict(zip(source_headers, values))
        record = OrderedDict()
        for target_key in OUTPUT_FIELD_ORDER:
            source_key = next(key for key, value in HEADER_MAP.items() if value == target_key)
            raw_value = source_row.get(source_key)
            if target_key == "data":
                record[target_key] = normalize_date(raw_value)
            elif target_key in {"num_viagens", "volume_m3", "valor"}:
                record[target_key] = normalize_number(raw_value)
            else:
                record[target_key] = normalize_text(raw_value)
        rows.append(record)

    return rows


def build_payload(source_path, rows):
    fields = OUTPUT_FIELD_ORDER
    return OrderedDict([
        ("meta", OrderedDict([
            ("gerado_em", datetime.now(TZ).isoformat(timespec="seconds")),
            ("fonte", os.path.basename(source_path)),
            ("planilha", SOURCE_SHEET),
            ("total_registros", len(rows)),
            ("campos", fields),
        ])),
        ("dados", rows),
    ])


def append_rows_to_existing_json(source_path, start_row, end_row):
    with open(OUTPUT_FILE, "r", encoding="utf-8") as fh:
        payload = json.load(fh)

    new_rows = read_rows_by_excel_range(source_path, start_row, end_row)
    existing = payload.get("dados", [])

    existing_keys = {json.dumps(row, ensure_ascii=False, sort_keys=True) for row in existing}
    appended = []
    for row in new_rows:
        key = json.dumps(row, ensure_ascii=False, sort_keys=True)
        if key in existing_keys:
            continue
        existing.append(row)
        existing_keys.add(key)
        appended.append(row)

    payload["meta"]["gerado_em"] = datetime.now(TZ).isoformat(timespec="seconds")
    payload["meta"]["fonte"] = os.path.basename(source_path)
    payload["meta"]["planilha"] = SOURCE_SHEET
    payload["meta"]["campos"] = OUTPUT_FIELD_ORDER
    payload["meta"]["total_registros"] = len(existing)

    with open(OUTPUT_FILE, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2)

    return {
        "source": source_path,
        "output": OUTPUT_FILE,
        "sheet": SOURCE_SHEET,
        "excel_range": [start_row, end_row],
        "rows_requested": len(new_rows),
        "rows_appended": len(appended),
        "total_registros": len(existing),
        "first_appended": appended[0] if appended else None,
        "last_appended": appended[-1] if appended else None,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source")
    parser.add_argument("--append-start-row", type=int)
    parser.add_argument("--append-end-row", type=int)
    args = parser.parse_args()

    source_path = args.source or find_source_file()
    if args.append_start_row and args.append_end_row:
        result = append_rows_to_existing_json(
            source_path,
            args.append_start_row,
            args.append_end_row,
        )
    else:
        rows = read_rows(source_path)
        payload = build_payload(source_path, rows)

        with open(OUTPUT_FILE, "w", encoding="utf-8") as fh:
            json.dump(payload, fh, ensure_ascii=False, indent=2)

        result = {
            "source": source_path,
            "output": OUTPUT_FILE,
            "sheet": SOURCE_SHEET,
            "total_registros": len(rows),
            "first_record": rows[0] if rows else None,
            "last_record": rows[-1] if rows else None,
        }

    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
